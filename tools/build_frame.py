#!/usr/bin/env python3
"""
DART 공시목록에서 2016년 이후 인선이엔티 정기보고서를 골라
'분기 ↔ 출처 보고서' 매핑표(수집 프레임)와 출처 시트를 만든다.

원문 수집 전에는 단가 열이 비어 있고, tools/dart_fetch.py + tools/dart_extract.py
를 돌리면 같은 구조에 값이 채워진다.
"""
from __future__ import annotations

import argparse
import re
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

QMAP = {"03": "1Q", "06": "2Q", "09": "3Q", "12": "4Q"}
BASIS = {"1Q": "1분기 누계(1~3월)", "2Q": "반기 누계(1~6월)",
         "3Q": "3분기 누계(1~9월)", "4Q": "연간 누계(1~12월)"}
PERIODIC = re.compile(r"(사업보고서|반기보고서|분기보고서)")
AUDIT = re.compile(r"감사보고서")

# 정기보고서 II.사업의 내용에서 단가가 실리는 위치
SECTIONS = [
    ("매립 단가", "II. 사업의 내용 > 2. 주요 제품 및 서비스 > 주요 제품 등의 가격변동추이",
     "매립폐기물 최종처리용역 단가 (원/톤). 매립 매출액 ÷ 처리량으로 산출"),
    ("소각 단가", "II. 사업의 내용 > 2. 주요 제품 및 서비스 > 주요 제품 등의 가격변동추이",
     "소각 처리용역 단가 (원/톤)"),
    ("건설폐기물 중간처리 단가", "II. 사업의 내용 > 2. 주요 제품 및 서비스 > 주요 제품 등의 가격변동추이",
     "건설폐기물 중간처리 단가 (원/톤)"),
    ("순환골재 판매단가", "II. 사업의 내용 > 2. 주요 제품 및 서비스 > 주요 제품 등의 가격변동추이",
     "순환골재 판매단가 (원/톤)"),
    ("자동차재활용(폐차) 단가", "II. 사업의 내용 > 2. 주요 제품 및 서비스 > 주요 제품 등의 가격변동추이",
     "폐차 대당 단가 및 고철/비철 판매단가"),
    ("수집·운반 단가", "II. 사업의 내용 > 2. 주요 제품 및 서비스 > 주요 제품 등의 가격변동추이",
     "폐기물 수집운반 용역 단가"),
    ("주요 원재료 매입단가", "II. 사업의 내용 > 3. 원재료 및 생산설비 > 주요 원재료 가격변동추이",
     "경유 등 주요 원재료·유틸리티 매입단가"),
    ("매출 실적 단가 환산", "II. 사업의 내용 > 4. 매출 및 수주상황",
     "품목별 매출액·물량으로 단가 역산 검증용"),
]

HDR_FILL = PatternFill("solid", fgColor="D6D2C4")
NOTE_FILL = PatternFill("solid", fgColor="F7EFDA")
THIN = Side(style="thin", color="A5AAAA")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def add_sheet(wb, title, headers, rows, widths=None, wrap_cols=()):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(1, c)
        cell.fill, cell.border = HDR_FILL, BORDER
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r in rows:
        ws.append(r)
    for c in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(c)].width = (widths or {}).get(c, 16)
        if c in wrap_cols:
            for r in range(2, ws.max_row + 1):
                ws.cell(r, c).alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
    return ws


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", required=True, type=Path)
    ap.add_argument("--out", required=True, type=Path)
    ap.add_argument("--since", default="2016-01-01")
    args = ap.parse_args()

    src = openpyxl.load_workbook(args.xlsx, data_only=True)
    ws = src["공시목록"]
    head = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    i = {n: head.index(n) for n in ("접수일자", "보고서명", "제출인", "접수번호", "공시 링크")}

    periodic, audit = [], []
    for row in ws.iter_rows(min_row=2, values_only=True):
        name = (row[i["보고서명"]] or "").strip()
        date = row[i["접수일자"]] or ""
        rec = {"date": date, "name": name, "filer": row[i["제출인"]],
               "rcp": str(row[i["접수번호"]]).strip(), "url": row[i["공시 링크"]]}
        if date < args.since:
            continue
        if PERIODIC.search(name):
            m = re.search(r"\((\d{4})\.(\d{2})\)", name)
            if m and m.group(2) in QMAP:
                rec["y"], rec["q"] = int(m.group(1)), QMAP[m.group(2)]
                periodic.append(rec)
        elif AUDIT.search(name):
            audit.append(rec)
    periodic.sort(key=lambda r: r["date"])
    audit.sort(key=lambda r: r["date"])

    wb = openpyxl.Workbook()

    # 1) 매립단가 분기 프레임
    qorder = {"1Q": 1, "2Q": 2, "3Q": 3, "4Q": 4}
    rows = [[r["y"], r["q"], f"{r['y']} {r['q']}", "매립", None, "원/톤",
             BASIS[r["q"]], r["name"], r["date"], r["rcp"], r["url"], "미수집"]
            for r in sorted(periodic, key=lambda r: (r["y"], qorder[r["q"]]))]
    add_sheet(wb, "매립단가_분기",
              ["연도", "분기", "기간", "항목", "단가", "단위", "집계기준",
               "출처 보고서", "접수일자", "접수번호", "DART 원문 링크", "수집상태"],
              rows, widths={3: 12, 5: 14, 7: 20, 8: 24, 11: 50, 12: 12})

    # 2) 발췌 대상 단가 항목 정의
    add_sheet(wb, "발췌대상_단가항목",
              ["단가 항목", "보고서 내 위치", "정의 / 비고"],
              [[a, b, c] for a, b, c in SECTIONS],
              widths={1: 26, 2: 62, 3: 54}, wrap_cols=(2, 3))

    # 3) 출처 공시목록 (정기보고서)
    add_sheet(wb, "출처_정기보고서",
              ["접수일자", "보고서명", "대상 분기", "제출인", "접수번호", "DART 원문 링크"],
              [[r["date"], r["name"], f"{r['y']} {r['q']}", r["filer"], r["rcp"], r["url"]]
               for r in periodic],
              widths={2: 28, 3: 12, 6: 50})

    # 4) 출처 감사보고서
    add_sheet(wb, "출처_감사보고서",
              ["접수일자", "보고서명", "제출인", "접수번호", "DART 원문 링크"],
              [[r["date"], r["name"], r["filer"], r["rcp"], r["url"]] for r in audit],
              widths={2: 28, 5: 50})

    wb.remove(wb["Sheet"])
    args.out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.out)
    print(f"저장: {args.out}")
    print(f"  정기보고서 {len(periodic)}건 ({periodic[0]['y']} {periodic[0]['q']} "
          f"~ {periodic[-1]['y']} {periodic[-1]['q']}), 감사보고서 {len(audit)}건")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
