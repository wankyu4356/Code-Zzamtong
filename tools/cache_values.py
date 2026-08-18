#!/usr/bin/env python3
"""수식 셀에 계산 결과를 캐시값으로 주입한다.

openpyxl 은 수식을 문자열로만 쓰고 캐시값을 남기지 않아, 재계산하지 않는 뷰어에서는
수식 셀이 빈 칸으로 보인다. LibreOffice 가 이 환경에서 xlsx 를 못 열기 때문에
formulas 라이브러리로 계산한 값을 시트 XML 의 <f> 뒤에 <v> 로 넣어준다.
수식 자체는 그대로 남으므로 엑셀에서 열면 정상적으로 재계산된다.
"""
from __future__ import annotations

import re
import shutil
import sys
import warnings
import zipfile
from pathlib import Path

warnings.filterwarnings("ignore")

# 빈 셀은 <c r="A1" s="2"/> 처럼 자기닫힘으로 쓰인다. (?<!/) 로 그런 태그를 제외하지
# 않으면 매치가 다음 셀들까지 삼켜 수식 셀을 놓친다.
CELL = re.compile(
    rb'<c\b[^>]*\br="([A-Z]+\d+)"[^>]*(?<!/)>(?:(?!</c>).)*?</c>', re.S)
HAS_F = re.compile(rb"<f[ >]")
# openpyxl 은 수식 셀에 빈 <v></v> 를 함께 쓴다. 그 자리를 계산값으로 채운다.
EMPTY_V = re.compile(rb"<v\s*/>|<v></v>")
FILLED_V = re.compile(rb"<v>[^<]+</v>")


def esc(s: str) -> bytes:
    return (s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            ).encode("utf-8")


def compute(path: Path) -> dict[tuple[str, str], object]:
    import formulas
    sol = formulas.ExcelModel().loads(str(path)).finish().calculate()
    out = {}
    for key, val in sol.items():
        if "!" not in key or "<" in key:
            continue
        sheet, _, coord = key.rpartition("!")
        sheet = sheet.strip("'").split("]")[-1]
        try:
            out[(sheet, coord)] = val.value[0, 0]
        except Exception:  # noqa: BLE001
            pass
    return out


def main() -> int:
    src = Path(sys.argv[1])
    values = compute(src)

    import openpyxl
    wb = openpyxl.load_workbook(src)
    # 시트 이름 -> xl/worksheets/sheetN.xml 순서 매핑
    order = {i + 1: name for i, name in enumerate(wb.sheetnames)}
    wb.close()

    tmp = src.with_suffix(".tmp.xlsx")
    injected = 0
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(
            tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            m = re.fullmatch(r"xl/worksheets/sheet(\d+)\.xml", item.filename)
            if m and int(m.group(1)) in order:
                name = order[int(m.group(1))]

                def repl(mo: re.Match) -> bytes:
                    nonlocal injected
                    cell = mo.group(0)
                    if not HAS_F.search(cell) or FILLED_V.search(cell):
                        return cell
                    v = values.get((name, mo.group(1).decode()))
                    if v is None or isinstance(v, bool):
                        return cell
                    if isinstance(v, (int, float)):
                        body = b"<v>%s</v>" % repr(float(v)).encode()
                    else:
                        text = str(v)
                        if text.strip() in ("", "[]"):
                            return cell      # 빈 문자열은 캐시하지 않는다
                        cell = cell.replace(b"<c ", b'<c t="str" ', 1)
                        body = b"<v>" + esc(text) + b"</v>"
                    injected += 1
                    if EMPTY_V.search(cell):
                        return EMPTY_V.sub(body, cell, count=1)
                    return cell.replace(b"</c>", body + b"</c>", 1)

                data = CELL.sub(repl, data)
            zout.writestr(item, data)

    shutil.move(tmp, src)
    print(f"캐시값 주입: {injected}셀")

    # 주입 후에도 파일이 정상인지 확인
    wb2 = openpyxl.load_workbook(src, data_only=True)
    ok = sum(1 for ws in wb2 for row in ws.iter_rows() for c in row
             if c.value is not None)
    print(f"검증: 파일 정상 로드, 값이 있는 셀 {ok:,}개")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
