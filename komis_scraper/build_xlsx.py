# -*- coding: utf-8 -*-
"""scrape.py가 저장한 JSON(data/)과 코드표(raw/)로 엑셀 보고서를 생성. 생성 후 LibreOffice로 재계산 권장."""
import json, os, glob, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment

HERE = os.path.dirname(os.path.abspath(__file__))
D = os.environ.get("KOMIS_DATA_DIR", os.path.join(HERE, "data"))
RAW = os.environ.get("KOMIS_RAW_DIR", os.path.join(HERE, "raw"))
OUT = os.environ.get("KOMIS_OUT", os.path.join(os.path.dirname(HERE), "KOMIS_희소금속_수입현황_2020-2026.xlsx"))

METALS = [  # code, symbol, korean, english
    ("MNRL0026", "Ta", "탄탈륨", "Tantalum"),
    ("MNRL0007", "Nb", "니오븀", "Niobium"),
    ("MNRL0025", "In", "인듐", "Indium"),
    ("MNRL0024", "Ga", "갈륨", "Gallium"),
    ("MNRL0040", "Hf", "하프늄", "Hafnium"),
    ("MNRL0042", "Re", "레늄", "Rhenium"),
]
SYM = {c: s for c, s, _, _ in METALS}
KOR = {c: k for c, _, k, _ in METALS}
YEARS = list(range(2020, 2027))
MONTHS = [f"{y}{m:02d}" for y in YEARS for m in range(1, 13) if (y, m) <= (2026, 9)]
MF_NM = {"001": "기초원료", "002": "기초금속", "003": "가공금속", "004": "스크랩"}
HS_KR = {}
HS_UN = {}
for c, *_ in METALS:
    HS_KR[c] = json.load(open(os.path.join(RAW, f"hs_{c}.json"), encoding="utf-8"))["data"]
    HS_UN[c] = json.load(open(os.path.join(RAW, f"hsun_{c}.json"), encoding="utf-8"))["data"]
HS_KR_NM = {h["hsCd"]: h["itemNm"] for c in HS_KR for h in HS_KR[c]}
HS_UN_NM = {h["hsCdUn"]: h["rmrk"] for c in HS_UN for h in HS_UN[c]}

def load(k):
    p = os.path.join(D, k + ".json")
    if not os.path.exists(p):
        return None
    return json.load(open(p, encoding="utf-8"))["resp"]

# ---------- English country names from map data ----------
ENG = {}
for f in glob.glob(os.path.join(D, "K_map_*.json")):
    r = json.load(open(f, encoding="utf-8"))["resp"]
    for row in r.get("data", {}).get("mapData", []) + r.get("data", {}).get("top5Data", []):
        if row.get("fromNtnCd") and row.get("fromEngNm"):
            ENG[row["fromNtnCd"]] = row["fromEngNm"]
        if row.get("toNtnCd") and row.get("toNtnEngNm"):
            ENG.setdefault(row["toNtnCd"], row["toNtnEngNm"])
NTN = {}
try:
    for r in json.load(open(os.path.join(RAW, "ntn_codes.json"), encoding="utf-8"))["data"]:
        NTN[r["cdKey"]] = (r["cdVal"], r.get("grpVal", ""))
except Exception:
    pass

# ---------- styles ----------
FONT = "Arial"
f_norm = Font(name=FONT, size=10)
f_bold = Font(name=FONT, size=10, bold=True)
f_head = Font(name=FONT, size=10, bold=True, color="FFFFFF")
f_title = Font(name=FONT, size=14, bold=True)
f_note = Font(name=FONT, size=9, italic=True, color="555555")
fill_head = PatternFill("solid", fgColor="1F4E78")
fill_sub = PatternFill("solid", fgColor="DDEBF7")
fill_formula = PatternFill("solid", fgColor="F2F2F2")
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
left = Alignment(horizontal="left", vertical="center", wrap_text=False)

FMT_INT = "#,##0"
FMT_1 = "#,##0.0"
FMT_2 = "#,##0.00"
FMT_PCT = "0.00%"

wb = Workbook()
wb.remove(wb.active)

def new_sheet(name, headers, widths=None, note=None, freeze="A2"):
    ws = wb.create_sheet(name)
    r0 = 1
    if note:
        ws.cell(row=1, column=1, value=note).font = f_note
        r0 = 2
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=r0, column=i, value=h)
        c.font = f_head; c.fill = fill_head; c.alignment = center; c.border = border
    ws.row_dimensions[r0].height = 32
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = f"A{r0+1}"
    return ws, r0 + 1

def put(ws, r, c, v, fmt=None, bold=False, fill=None):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font = f_bold if bold else f_norm
    if fmt: cell.number_format = fmt
    if fill: cell.fill = fill
    return cell

def finish(ws, header_row, last_row, last_col):
    if last_row >= header_row:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(last_col)}{last_row}"

def q(name):  # quoted sheet ref
    return "'" + name + "'"

# =====================================================================
# 1) 관세청_연도별_국가 (Korea customs, per country, per year)
# =====================================================================
S_KY = "관세청_연도별_국가"
hdr = ["키", "광종코드", "광종", "원소", "연도", "순위(금액)", "국가코드", "국가명", "국가명(영문)",
       "수입금액(USD)", "수입금액(천USD)", "수입중량(kg)", "수입중량(톤)", "수입금액 비중", "수입중량 비중",
       "수출금액(USD)", "수출금액(천USD)", "수출중량(kg)", "수출중량(톤)", "수출금액 비중", "수출중량 비중",
       "KOMIS 총수입금액(USD)", "KOMIS 총수입중량(kg)", "KOMIS 총수출금액(USD)", "KOMIS 총수출중량(kg)", "순위(중량)", "중량키"]
ws, r = new_sheet(S_KY, hdr, [14, 10, 8, 6, 6, 5, 7, 12, 22, 14, 13, 13, 11, 10, 10, 14, 13, 13, 11, 10, 10, 16, 16, 16, 16, 6, 14],
                  note="출처: KOMIS 핵심광물지도-대한민국 (관세청 수출입통계). 금액 USD, 중량 kg 원자료. 천USD/톤/비중 열은 수식. 2026년은 1~9월 누계.")
KY_FIRST = r
for code, sym, kor, eng in METALS:
    for y in YEARS:
        d = load(f"K_list_{code}_{y}")
        rows = d["list"] if d else []
        wrank = {id(x): i for i, x in enumerate(sorted(rows, key=lambda x: (-x["incmWeig"], -x["incmAmt"])), 1)}
        for row in rows:
            cc = row["ntnCd"]
            put(ws, r, 1, f"{sym}|{y}|{row['RNUM']}")
            put(ws, r, 2, code); put(ws, r, 3, kor); put(ws, r, 4, sym); put(ws, r, 5, y); put(ws, r, 6, row["RNUM"])
            put(ws, r, 7, cc); put(ws, r, 8, row.get("ntnKornNm") or NTN.get(cc, ("",))[0]); put(ws, r, 9, ENG.get(cc, ""))
            put(ws, r, 10, row["incmAmt"], FMT_INT)
            put(ws, r, 11, f"=J{r}/1000", FMT_1, fill=fill_formula)
            put(ws, r, 12, row["incmWeig"], FMT_INT)
            put(ws, r, 13, f"=L{r}/1000", FMT_2, fill=fill_formula)
            put(ws, r, 14, f'=IF(SUMIFS($J$3:$J$20000,$D$3:$D$20000,$D{r},$E$3:$E$20000,$E{r})=0,0,J{r}/SUMIFS($J$3:$J$20000,$D$3:$D$20000,$D{r},$E$3:$E$20000,$E{r}))', FMT_PCT, fill=fill_formula)
            put(ws, r, 15, f'=IF(SUMIFS($L$3:$L$20000,$D$3:$D$20000,$D{r},$E$3:$E$20000,$E{r})=0,0,L{r}/SUMIFS($L$3:$L$20000,$D$3:$D$20000,$D{r},$E$3:$E$20000,$E{r}))', FMT_PCT, fill=fill_formula)
            put(ws, r, 16, row["expAmt"], FMT_INT)
            put(ws, r, 17, f"=P{r}/1000", FMT_1, fill=fill_formula)
            put(ws, r, 18, row["expWeig"], FMT_INT)
            put(ws, r, 19, f"=R{r}/1000", FMT_2, fill=fill_formula)
            put(ws, r, 20, f'=IF(SUMIFS($P$3:$P$20000,$D$3:$D$20000,$D{r},$E$3:$E$20000,$E{r})=0,0,P{r}/SUMIFS($P$3:$P$20000,$D$3:$D$20000,$D{r},$E$3:$E$20000,$E{r}))', FMT_PCT, fill=fill_formula)
            put(ws, r, 21, f'=IF(SUMIFS($R$3:$R$20000,$D$3:$D$20000,$D{r},$E$3:$E$20000,$E{r})=0,0,R{r}/SUMIFS($R$3:$R$20000,$D$3:$D$20000,$D{r},$E$3:$E$20000,$E{r}))', FMT_PCT, fill=fill_formula)
            put(ws, r, 22, row["sumIncmAmt"], FMT_INT); put(ws, r, 23, row["sumIncmWeig"], FMT_INT)
            put(ws, r, 24, row["sumExpAmt"], FMT_INT); put(ws, r, 25, row["sumExpWeig"], FMT_INT)
            put(ws, r, 26, wrank[id(row)]); put(ws, r, 27, f"{sym}|{y}|W|{wrank[id(row)]}")
            r += 1
KY_LAST = r - 1
finish(ws, KY_FIRST - 1, KY_LAST, len(hdr))

# =====================================================================
# 2) 요약 (summary with formulas)
# =====================================================================
S_SUM = "요약"
hdr = ["광종", "원소", "연도", "수입금액(천USD)", "수입중량(톤)", "수입금액 전년비", "수입중량 전년비",
       "수출금액(천USD)", "수출중량(톤)", "수입국 수",
       "1위 수입국", "1위 비중(금액)", "2위 수입국", "2위 비중(금액)", "3위 수입국", "3위 비중(금액)",
       "주요 수입국 Top3 (금액비중)", "1위 수입국(중량기준)", "1위 비중(중량)", "비고"]
ws, r = new_sheet(S_SUM, hdr, [8, 6, 6, 15, 13, 11, 11, 15, 13, 8, 12, 11, 12, 11, 12, 11, 44, 14, 11, 22],
                  note="관세청(KOMIS 대한민국 광물지도) 기준. 모든 수치는 '관세청_연도별_국가' 시트에서 수식으로 집계. 비중=해당국 수입금액/총수입금액. 2026년은 1~9월 누계(전년비는 연간 대비이므로 참고용).")
SUM_FIRST = r
KYq = q(S_KY)
for code, sym, kor, eng in METALS:
    for y in YEARS:
        put(ws, r, 1, kor); put(ws, r, 2, sym); put(ws, r, 3, y)
        put(ws, r, 4, f"=SUMIFS({KYq}!$J:$J,{KYq}!$D:$D,$B{r},{KYq}!$E:$E,$C{r})/1000", FMT_1, fill=fill_formula)
        put(ws, r, 5, f"=SUMIFS({KYq}!$L:$L,{KYq}!$D:$D,$B{r},{KYq}!$E:$E,$C{r})/1000", FMT_2, fill=fill_formula)
        if y == YEARS[0]:
            put(ws, r, 6, ""); put(ws, r, 7, "")
        else:
            put(ws, r, 6, f'=IF(D{r-1}=0,"",D{r}/D{r-1}-1)', FMT_PCT, fill=fill_formula)
            put(ws, r, 7, f'=IF(E{r-1}=0,"",E{r}/E{r-1}-1)', FMT_PCT, fill=fill_formula)
        put(ws, r, 8, f"=SUMIFS({KYq}!$P:$P,{KYq}!$D:$D,$B{r},{KYq}!$E:$E,$C{r})/1000", FMT_1, fill=fill_formula)
        put(ws, r, 9, f"=SUMIFS({KYq}!$R:$R,{KYq}!$D:$D,$B{r},{KYq}!$E:$E,$C{r})/1000", FMT_2, fill=fill_formula)
        put(ws, r, 10, f'=COUNTIFS({KYq}!$D:$D,$B{r},{KYq}!$E:$E,$C{r},{KYq}!$J:$J,">0")', FMT_INT, fill=fill_formula)
        col = 11
        for rank in (1, 2, 3):
            key = f'$B{r}&"|"&$C{r}&"|{rank}"'
            put(ws, r, col, f'=IFERROR(INDEX({KYq}!$H:$H,MATCH({key},{KYq}!$A:$A,0)),"")', fill=fill_formula)
            put(ws, r, col + 1, f'=IFERROR(INDEX({KYq}!$N:$N,MATCH({key},{KYq}!$A:$A,0)),"")', FMT_PCT, fill=fill_formula)
            col += 2
        put(ws, r, 17, f'=IF(K{r}="","",K{r}&" ("&TEXT(L{r},"0.0%")&")")&IF(M{r}="","",", "&M{r}&" ("&TEXT(N{r},"0.0%")&")")&IF(O{r}="","",", "&O{r}&" ("&TEXT(P{r},"0.0%")&")")', fill=fill_formula)
        # 1위 (중량기준): country with max import weight for this metal-year
        wkey = f'$B{r}&"|"&$C{r}&"|W|1"'
        put(ws, r, 18, f'=IFERROR(INDEX({KYq}!$H:$H,MATCH({wkey},{KYq}!$AA:$AA,0)),"")', fill=fill_formula)
        put(ws, r, 19, f'=IFERROR(INDEX({KYq}!$O:$O,MATCH({wkey},{KYq}!$AA:$AA,0)),"")', FMT_PCT, fill=fill_formula)
        put(ws, r, 20, "2026년은 1~9월 누계" if y == 2026 else "")
        r += 1
SUM_LAST = r - 1
finish(ws, SUM_FIRST - 1, SUM_LAST, len(hdr))

# =====================================================================
# 3) 요약_매트릭스 (metal x year matrices)
# =====================================================================
S_MAT = "요약_매트릭스"
ws = wb.create_sheet(S_MAT)
ws.cell(row=1, column=1, value="관세청 기준 광종별 연도별 수입 현황 (수식: '요약' 시트 참조). 2026년은 1~9월 누계.").font = f_note
SUMq = q(S_SUM)
blocks = [("수입금액(천USD)", "D", FMT_1), ("수입중량(톤)", "E", FMT_2), ("수출금액(천USD)", "H", FMT_1), ("수출중량(톤)", "I", FMT_2)]
rr = 3
block_rows = {}
for title, colL, fmt in blocks:
    ws.cell(row=rr, column=1, value=title).font = f_bold
    rr += 1
    put(ws, rr, 1, "광종", bold=True, fill=fill_sub); put(ws, rr, 2, "원소", bold=True, fill=fill_sub)
    for j, y in enumerate(YEARS):
        c = put(ws, rr, 3 + j, y, bold=True, fill=fill_sub); c.alignment = center
    hdr_row = rr; rr += 1
    block_rows[title] = {}
    for code, sym, kor, eng in METALS:
        put(ws, rr, 1, kor); put(ws, rr, 2, sym)
        block_rows[title][sym] = rr
        for j, y in enumerate(YEARS):
            cl = get_column_letter(3 + j)
            put(ws, rr, 3 + j, f'=SUMIFS({SUMq}!${colL}:${colL},{SUMq}!$B:$B,$B{rr},{SUMq}!$C:$C,{cl}${hdr_row})', fmt, fill=fill_formula)
        rr += 1
    rr += 1
for title, src in [("수입금액 전년비(%)", "수입금액(천USD)"), ("수입중량 전년비(%)", "수입중량(톤)")]:
    ws.cell(row=rr, column=1, value=title + "  (2026년은 1~9월 누계 대비이므로 참고용)").font = f_bold
    rr += 1
    put(ws, rr, 1, "광종", bold=True, fill=fill_sub); put(ws, rr, 2, "원소", bold=True, fill=fill_sub)
    for j, y in enumerate(YEARS):
        c = put(ws, rr, 3 + j, y, bold=True, fill=fill_sub); c.alignment = center
    rr += 1
    for code, sym, kor, eng in METALS:
        put(ws, rr, 1, kor); put(ws, rr, 2, sym)
        sr = block_rows[src][sym]
        for j, y in enumerate(YEARS):
            if j == 0:
                put(ws, rr, 3, "")
            else:
                cur = f"{get_column_letter(3+j)}{sr}"; prv = f"{get_column_letter(2+j)}{sr}"
                put(ws, rr, 3 + j, f'=IF({prv}=0,"",{cur}/{prv}-1)', FMT_PCT, fill=fill_formula)
        rr += 1
    rr += 1
ws.column_dimensions["A"].width = 10; ws.column_dimensions["B"].width = 6
for j in range(len(YEARS)):
    ws.column_dimensions[get_column_letter(3 + j)].width = 14

# =====================================================================
# 4) 관세청_추이(API) : line chart data
# =====================================================================
S_LINE = "관세청_연도별_총계(API)"
hdr = ["광종코드", "광종", "원소", "연도", "총수입금액(USD)", "총수입금액(천USD)", "총수입중량(kg)", "총수입중량(톤)",
       "총수출금액(USD)", "총수출금액(천USD)", "총수출중량(kg)", "총수출중량(톤)"]
ws, r = new_sheet(S_LINE, hdr, [10, 8, 6, 6, 16, 15, 15, 13, 16, 15, 15, 13],
                  note="출처: KOMIS 대한민국 광물지도 '수입규모 추이' 차트 API(getLineChartDataKorea). '요약' 시트 SUMIFS 결과와 교차검증용. 2026년은 1~9월 누계.")
LINE_FIRST = r
for code, sym, kor, eng in METALS:
    merged = {}
    for ye in (2024, 2026):
        d = load(f"K_line_{code}_{ye}")
        for row in (d or {}).get("data", []):
            merged[int(row["crtrYmd"])] = row
    for y in sorted(merged):
        if y not in YEARS: continue
        row = merged[y]
        put(ws, r, 1, code); put(ws, r, 2, kor); put(ws, r, 3, sym); put(ws, r, 4, y)
        put(ws, r, 5, row["totalIncmAmt"], FMT_INT); put(ws, r, 6, f"=E{r}/1000", FMT_1, fill=fill_formula)
        put(ws, r, 7, row["totalIncmWeig"], FMT_INT); put(ws, r, 8, f"=G{r}/1000", FMT_2, fill=fill_formula)
        put(ws, r, 9, row["totalExpAmt"], FMT_INT); put(ws, r, 10, f"=I{r}/1000", FMT_1, fill=fill_formula)
        put(ws, r, 11, row["totalExpWeig"], FMT_INT); put(ws, r, 12, f"=K{r}/1000", FMT_2, fill=fill_formula)
        r += 1
finish(ws, LINE_FIRST - 1, r - 1, len(hdr))

# =====================================================================
# 5) 관세청_HS코드별 (per HS10, per country, per year)
# =====================================================================
S_HS = "관세청_HS코드별_국가"
hdr = ["광종", "원소", "연도", "HS코드(10자리)", "품목명", "순위", "국가코드", "국가명",
       "수입금액(USD)", "수입금액(천USD)", "수입중량(kg)", "수입중량(톤)", "수입금액 비중(HS내)", "수입중량 비중(HS내)",
       "수출금액(USD)", "수출금액(천USD)", "수출중량(kg)", "수출중량(톤)", "HS 총수입금액(USD)", "HS 총수입중량(kg)"]
ws, r = new_sheet(S_HS, hdr, [8, 6, 6, 14, 34, 5, 7, 12, 14, 13, 13, 11, 12, 12, 14, 13, 13, 11, 16, 16],
                  note="출처: KOMIS 대한민국 광물지도, HS코드(관세청 10자리)별 국가별 수출입. 비중은 해당 HS코드 내 수식. 2026년은 1~9월 누계.")
HS_FIRST = r
for code, sym, kor, eng in METALS:
    for y in YEARS:
        for h in HS_KR[code]:
            hs = h["hsCd"]
            d = load(f"K_lisths_{code}_{y}_{hs}")
            for row in (d or {}).get("list", []):
                put(ws, r, 1, kor); put(ws, r, 2, sym); put(ws, r, 3, y); put(ws, r, 4, hs); put(ws, r, 5, row.get("itemNm") or HS_KR_NM.get(hs, ""))
                put(ws, r, 6, row["RNUM"]); put(ws, r, 7, row["ntnCd"]); put(ws, r, 8, row["ntnKornNm"])
                put(ws, r, 9, row["incmAmt"], FMT_INT); put(ws, r, 10, f"=I{r}/1000", FMT_1, fill=fill_formula)
                put(ws, r, 11, row["incmWeig"], FMT_INT); put(ws, r, 12, f"=K{r}/1000", FMT_2, fill=fill_formula)
                put(ws, r, 13, f'=IF(SUMIFS($I$3:$I$30000,$B$3:$B$30000,$B{r},$C$3:$C$30000,$C{r},$D$3:$D$30000,$D{r})=0,0,I{r}/SUMIFS($I$3:$I$30000,$B$3:$B$30000,$B{r},$C$3:$C$30000,$C{r},$D$3:$D$30000,$D{r}))', FMT_PCT, fill=fill_formula)
                put(ws, r, 14, f'=IF(SUMIFS($K$3:$K$30000,$B$3:$B$30000,$B{r},$C$3:$C$30000,$C{r},$D$3:$D$30000,$D{r})=0,0,K{r}/SUMIFS($K$3:$K$30000,$B$3:$B$30000,$B{r},$C$3:$C$30000,$C{r},$D$3:$D$30000,$D{r}))', FMT_PCT, fill=fill_formula)
                put(ws, r, 15, row["expAmt"], FMT_INT); put(ws, r, 16, f"=O{r}/1000", FMT_1, fill=fill_formula)
                put(ws, r, 17, row["expWeig"], FMT_INT); put(ws, r, 18, f"=Q{r}/1000", FMT_2, fill=fill_formula)
                put(ws, r, 19, row["sumIncmAmt"], FMT_INT); put(ws, r, 20, row["sumIncmWeig"], FMT_INT)
                r += 1
HS_LAST = r - 1
finish(ws, HS_FIRST - 1, HS_LAST, len(hdr))

# 5b) HS code totals matrix
S_HST = "관세청_HS코드별_총계"
ws = wb.create_sheet(S_HST)
ws.cell(row=1, column=1, value="관세청 HS코드(10자리)별 연도별 총수입 (수식: '관세청_HS코드별_국가' 시트 SUMIFS). 2026년은 1~9월 누계.").font = f_note
HSq = q(S_HS)
rr = 3
for title, colL, fmt in [("수입금액(천USD)", "I", FMT_1), ("수입중량(톤)", "K", FMT_2), ("수출금액(천USD)", "O", FMT_1), ("수출중량(톤)", "Q", FMT_2)]:
    ws.cell(row=rr, column=1, value=title).font = f_bold; rr += 1
    for i, h in enumerate(["광종", "원소", "HS코드", "품목명"] + YEARS, 1):
        c = put(ws, rr, i, h, bold=True, fill=fill_sub); c.alignment = center
    hdr_row = rr; rr += 1
    for code, sym, kor, eng in METALS:
        for h in HS_KR[code]:
            put(ws, rr, 1, kor); put(ws, rr, 2, sym); put(ws, rr, 3, h["hsCd"]); put(ws, rr, 4, h["itemNm"])
            for j, y in enumerate(YEARS):
                cl = get_column_letter(5 + j)
                put(ws, rr, 5 + j, f"=SUMIFS({HSq}!${colL}:${colL},{HSq}!$B:$B,$B{rr},{HSq}!$D:$D,$C{rr},{HSq}!$C:$C,{cl}${hdr_row})/1000", fmt, fill=fill_formula)
            rr += 1
    rr += 1
for i, w in enumerate([8, 6, 13, 36] + [13] * len(YEARS), 1):
    ws.column_dimensions[get_column_letter(i)].width = w

# =====================================================================
# 6) 관세청_생산품유형별
# =====================================================================
S_MF = "관세청_생산품유형별_국가"
hdr = ["광종", "원소", "연도", "유형코드", "생산품 유형", "순위", "국가코드", "국가명",
       "수입금액(USD)", "수입금액(천USD)", "수입중량(kg)", "수입중량(톤)", "수입금액 비중(유형내)",
       "수출금액(USD)", "수출금액(천USD)", "수출중량(kg)", "수출중량(톤)", "유형 총수입금액(USD)", "유형 총수입중량(kg)"]
ws, r = new_sheet(S_MF, hdr, [8, 6, 6, 8, 10, 5, 7, 12, 14, 13, 13, 11, 12, 14, 13, 13, 11, 16, 16],
                  note="출처: KOMIS 대한민국 광물지도, 생산품 유형(물질흐름: 기초원료/기초금속/가공금속/스크랩)별 국가별 수출입. 2026년은 1~9월 누계.")
MF_FIRST = r
for code, sym, kor, eng in METALS:
    for y in YEARS:
        for mf in MF_NM:
            d = load(f"K_listmf_{code}_{y}_{mf}")
            for row in (d or {}).get("list", []):
                put(ws, r, 1, kor); put(ws, r, 2, sym); put(ws, r, 3, y); put(ws, r, 4, mf); put(ws, r, 5, MF_NM[mf])
                put(ws, r, 6, row["RNUM"]); put(ws, r, 7, row["ntnCd"]); put(ws, r, 8, row["ntnKornNm"])
                put(ws, r, 9, row["incmAmt"], FMT_INT); put(ws, r, 10, f"=I{r}/1000", FMT_1, fill=fill_formula)
                put(ws, r, 11, row["incmWeig"], FMT_INT); put(ws, r, 12, f"=K{r}/1000", FMT_2, fill=fill_formula)
                put(ws, r, 13, f'=IF(SUMIFS($I$3:$I$30000,$B$3:$B$30000,$B{r},$C$3:$C$30000,$C{r},$D$3:$D$30000,$D{r})=0,0,I{r}/SUMIFS($I$3:$I$30000,$B$3:$B$30000,$B{r},$C$3:$C$30000,$C{r},$D$3:$D$30000,$D{r}))', FMT_PCT, fill=fill_formula)
                put(ws, r, 14, row["expAmt"], FMT_INT); put(ws, r, 15, f"=N{r}/1000", FMT_1, fill=fill_formula)
                put(ws, r, 16, row["expWeig"], FMT_INT); put(ws, r, 17, f"=P{r}/1000", FMT_2, fill=fill_formula)
                put(ws, r, 18, row["sumIncmAmt"], FMT_INT); put(ws, r, 19, row["sumIncmWeig"], FMT_INT)
                r += 1
finish(ws, MF_FIRST - 1, r - 1, len(hdr))

S_MFT = "관세청_생산품유형별_총계"
ws = wb.create_sheet(S_MFT)
ws.cell(row=1, column=1, value="관세청 생산품 유형별 연도별 총수입 (수식: '관세청_생산품유형별_국가' SUMIFS). 2026년은 1~9월 누계.").font = f_note
MFq = q(S_MF)
rr = 3
for title, colL, fmt in [("수입금액(천USD)", "I", FMT_1), ("수입중량(톤)", "K", FMT_2), ("수출금액(천USD)", "N", FMT_1), ("수출중량(톤)", "P", FMT_2)]:
    ws.cell(row=rr, column=1, value=title).font = f_bold; rr += 1
    for i, h in enumerate(["광종", "원소", "유형코드", "생산품 유형"] + YEARS, 1):
        c = put(ws, rr, i, h, bold=True, fill=fill_sub); c.alignment = center
    hdr_row = rr; rr += 1
    for code, sym, kor, eng in METALS:
        for mf in MF_NM:
            put(ws, rr, 1, kor); put(ws, rr, 2, sym); put(ws, rr, 3, mf); put(ws, rr, 4, MF_NM[mf])
            for j, y in enumerate(YEARS):
                cl = get_column_letter(5 + j)
                put(ws, rr, 5 + j, f"=SUMIFS({MFq}!${colL}:${colL},{MFq}!$B:$B,$B{rr},{MFq}!$D:$D,$C{rr},{MFq}!$C:$C,{cl}${hdr_row})/1000", fmt, fill=fill_formula)
            rr += 1
    rr += 1
for i, w in enumerate([8, 6, 9, 12] + [13] * len(YEARS), 1):
    ws.column_dimensions[get_column_letter(i)].width = w

# =====================================================================
# 7) 관세청_월별_국가
# =====================================================================
S_MO = "관세청_월별_국가"
hdr = ["광종", "원소", "연월", "연도", "월", "순위", "국가코드", "국가명",
       "수입금액(USD)", "수입금액(천USD)", "수입중량(kg)", "수입중량(톤)", "수입금액 비중(당월)",
       "수출금액(USD)", "수출금액(천USD)", "수출중량(kg)", "수출중량(톤)", "당월 총수입금액(USD)", "당월 총수입중량(kg)"]
ws, r = new_sheet(S_MO, hdr, [8, 6, 9, 6, 4, 5, 7, 12, 14, 13, 13, 11, 12, 14, 13, 13, 11, 16, 16],
                  note="출처: KOMIS 대한민국 광물지도 월별 조회(2020.01~2026.09). 국가별 월별 수출입.")
MO_FIRST = r
for code, sym, kor, eng in METALS:
    for ym in MONTHS:
        d = load(f"K_month_{code}_{ym}")
        for row in (d or {}).get("list", []):
            put(ws, r, 1, kor); put(ws, r, 2, sym); put(ws, r, 3, f"{ym[:4]}-{ym[4:]}"); put(ws, r, 4, int(ym[:4])); put(ws, r, 5, int(ym[4:]))
            put(ws, r, 6, row["RNUM"]); put(ws, r, 7, row["ntnCd"]); put(ws, r, 8, row["ntnKornNm"])
            put(ws, r, 9, row["incmAmt"], FMT_INT); put(ws, r, 10, f"=I{r}/1000", FMT_1, fill=fill_formula)
            put(ws, r, 11, row["incmWeig"], FMT_INT); put(ws, r, 12, f"=K{r}/1000", FMT_2, fill=fill_formula)
            put(ws, r, 13, f'=IF(R{r}=0,0,I{r}/R{r})', FMT_PCT, fill=fill_formula)
            put(ws, r, 14, row["expAmt"], FMT_INT); put(ws, r, 15, f"=N{r}/1000", FMT_1, fill=fill_formula)
            put(ws, r, 16, row["expWeig"], FMT_INT); put(ws, r, 17, f"=P{r}/1000", FMT_2, fill=fill_formula)
            put(ws, r, 18, row["sumIncmAmt"], FMT_INT); put(ws, r, 19, row["sumIncmWeig"], FMT_INT)
            r += 1
MO_LAST = r - 1
finish(ws, MO_FIRST - 1, MO_LAST, len(hdr))

S_MOT = "관세청_월별_총계"
ws = wb.create_sheet(S_MOT)
ws.cell(row=1, column=1, value="관세청 월별 총수입 (수식: '관세청_월별_국가' SUMIFS). 좌: 수입금액(천USD), 우: 수입중량(톤).").font = f_note
MOq = q(S_MO)
hdrs = ["연월", "연도", "월"] + [f"{s} 수입금액(천USD)" for _, s, _, _ in METALS] + [f"{s} 수입중량(톤)" for _, s, _, _ in METALS]
for i, h in enumerate(hdrs, 1):
    c = put(ws, 3, i, h, bold=True, fill=fill_sub); c.alignment = center
rr = 4
for ym in MONTHS:
    put(ws, rr, 1, f"{ym[:4]}-{ym[4:]}"); put(ws, rr, 2, int(ym[:4])); put(ws, rr, 3, int(ym[4:]))
    for j, (code, sym, kor, eng) in enumerate(METALS):
        put(ws, rr, 4 + j, f'=SUMIFS({MOq}!$I:$I,{MOq}!$B:$B,"{sym}",{MOq}!$C:$C,$A{rr})/1000', FMT_1, fill=fill_formula)
        put(ws, rr, 4 + len(METALS) + j, f'=SUMIFS({MOq}!$K:$K,{MOq}!$B:$B,"{sym}",{MOq}!$C:$C,$A{rr})/1000', FMT_2, fill=fill_formula)
    rr += 1
ws.freeze_panes = "D4"
for i in range(1, len(hdrs) + 1):
    ws.column_dimensions[get_column_letter(i)].width = 10 if i <= 3 else 15

# =====================================================================
# 8) 관세청_전년대비증가TOP5
# =====================================================================
S_T5 = "관세청_전년대비증가TOP5"
hdr = ["광종", "원소", "순위", "국가코드", "국가명", "2021 수입금액(USD)", "2022 수입금액(USD)", "2023 수입금액(USD)", "2024 수입금액(USD)", "2025 수입금액(USD)", "2025년 전년대비 증가율(%)"]
ws, r = new_sheet(S_T5, hdr, [8, 6, 5, 8, 12, 16, 16, 16, 16, 16, 18],
                  note="출처: KOMIS 대한민국 광물지도 '전년대비 교역증가 TOP5' API(getTop5TableDataKorea). 조회연도와 무관하게 최근 5개년(2021~2025) 수입금액 기준으로 반환됨(검증: 프랑스 2024=75,556 / 2025=286,454 USD).")
T5_FIRST = r
for code, sym, kor, eng in METALS:
    d = load(f"K_top5_{code}_2025") or load(f"K_top5_{code}_2024")
    for i, row in enumerate((d or {}).get("data", []), 1):
        put(ws, r, 1, kor); put(ws, r, 2, sym); put(ws, r, 3, i); put(ws, r, 4, row["trgtNtnCd"]); put(ws, r, 5, row["ntnKornNm"])
        for j, k in enumerate(["tradeVolume_5", "tradeVolume_4", "tradeVolume_3", "tradeVolume_2", "tradeVolume_1"]):
            put(ws, r, 6 + j, row[k], FMT_INT)
        put(ws, r, 11, row["flutRateTradeVolume"], FMT_2)
        r += 1
finish(ws, T5_FIRST - 1, r - 1, len(hdr))

# =====================================================================
# 9) UN_수출국→한국_연도별
# =====================================================================
UN_NOTE = ("출처: KOMIS 핵심광물지도-글로벌 (UN Comtrade, HS 6자리). 한국 수입 기준 2025년 이후 미집계. "
           "주의: 인듐·갈륨·니오븀은 UN HS 6자리(811292/811299)가 Ga·Ge·In·Nb·V 공용이므로 금속별 구분 불가 → 인듐과 갈륨 수치가 동일하게 집계됨. "
           "탄탈륨·니오븀 광석(261590)도 Nb·Ta·V 공용. 금속별 정확한 수치는 관세청(10자리) 시트를 사용할 것.")
S_UN = "UN_수출국→한국_연도별"
hdr = ["광종", "원소", "연도", "순위", "수출국코드", "수출국명", "수출국명(영문)", "금액(USD)", "금액(천USD)", "중량(kg)", "중량(톤)",
       "금액 비중(%, KOMIS)", "중량 비중(%, KOMIS)", "한국 총수입금액(USD)", "한국 총수입중량(kg)", "비고"]
ws, r = new_sheet(S_UN, hdr, [8, 6, 6, 5, 8, 12, 22, 15, 13, 14, 12, 12, 12, 16, 16, 30], note=UN_NOTE)
UN_FIRST = r
for code, sym, kor, eng in METALS:
    for y in YEARS:
        d = load(f"N_list_krimp_{code}_{y}")
        rows = (d or {}).get("list", [])
        if not rows:
            put(ws, r, 1, kor); put(ws, r, 2, sym); put(ws, r, 3, y); put(ws, r, 16, "UN Comtrade 데이터 없음(미집계)")
            r += 1; continue
        for row in rows:
            put(ws, r, 1, kor); put(ws, r, 2, sym); put(ws, r, 3, y); put(ws, r, 4, row["RNUM"]); put(ws, r, 5, row["expNtnCd"]); put(ws, r, 6, row["expNtnNm"]); put(ws, r, 7, ENG.get(row["expNtnCd"], ""))
            put(ws, r, 8, row["amt"], FMT_INT); put(ws, r, 9, f"=H{r}/1000", FMT_1, fill=fill_formula)
            put(ws, r, 10, row["weig"], FMT_2); put(ws, r, 11, f"=J{r}/1000", FMT_2, fill=fill_formula)
            put(ws, r, 12, float(row["amtRate"]) / 100, FMT_PCT); put(ws, r, 13, float(row["weigRate"]) / 100, FMT_PCT)
            put(ws, r, 14, row["sumAmt"], FMT_INT); put(ws, r, 15, row["sumWeig"], FMT_2)
            put(ws, r, 16, "HS6 공용코드(Ga/Ge/In/Nb/V 합산)" if sym in ("In", "Ga", "Nb") else "")
            r += 1
finish(ws, UN_FIRST - 1, r - 1, len(hdr))

# 10) UN_HS6별
S_UNHS = "UN_HS6별_수출국→한국"
hdr = ["광종", "원소", "연도", "HS6", "HS6 품목(영문)", "순위", "수출국코드", "수출국명", "금액(USD)", "금액(천USD)", "중량(kg)", "중량(톤)", "금액 비중(%, KOMIS)", "중량 비중(%, KOMIS)", "HS6 총금액(USD)", "HS6 총중량(kg)"]
ws, r = new_sheet(S_UNHS, hdr, [8, 6, 6, 8, 40, 5, 8, 12, 15, 13, 14, 12, 12, 12, 16, 16], note=UN_NOTE)
UNHS_FIRST = r
for code, sym, kor, eng in METALS:
    for y in YEARS:
        for h in HS_UN[code]:
            hs = h["hsCdUn"]
            d = load(f"N_lisths_krimp_{code}_{y}_{hs}")
            for row in (d or {}).get("list", []):
                put(ws, r, 1, kor); put(ws, r, 2, sym); put(ws, r, 3, y); put(ws, r, 4, hs); put(ws, r, 5, row.get("RMRK") or HS_UN_NM.get(hs, ""))
                put(ws, r, 6, row["RNUM"]); put(ws, r, 7, row["expNtnCd"]); put(ws, r, 8, row["expNtnNm"])
                put(ws, r, 9, row["amt"], FMT_INT); put(ws, r, 10, f"=I{r}/1000", FMT_1, fill=fill_formula)
                put(ws, r, 11, row["weig"], FMT_2); put(ws, r, 12, f"=K{r}/1000", FMT_2, fill=fill_formula)
                put(ws, r, 13, float(row["amtRate"]) / 100, FMT_PCT); put(ws, r, 14, float(row["weigRate"]) / 100, FMT_PCT)
                put(ws, r, 15, row["sumAmt"], FMT_INT); put(ws, r, 16, row["sumWeig"], FMT_2)
                r += 1
finish(ws, UNHS_FIRST - 1, r - 1, len(hdr))

# 11) UN_수출국별_추이
S_UNB = "UN_수출국별_추이(한국수입)"
hdr = ["광종", "원소", "수출국코드", "수출국명"] + [f"{y} 금액(USD)" for y in YEARS] + ["합계(USD)"]
ws, r = new_sheet(S_UNB, hdr, [8, 6, 8, 12] + [15] * len(YEARS) + [16],
                  note="출처: KOMIS 글로벌 광물지도 '국가별 교역규모' 차트 API(getBarChartDataNation, 수입국=대한민국). UN Comtrade 기준, 2025년 이후 미집계(0). " + UN_NOTE.split("주의:")[1].strip())
UNB_FIRST = r
for code, sym, kor, eng in METALS:
    series = {}
    for ye in (2024, 2026):
        d = load(f"N_bar_krimp_{code}_{ye}")
        bc = (d or {}).get("data", {}).get("barChart", {})
        xs = bc.get("xaxis", [])
        for s in bc.get("series", []):
            ent = series.setdefault(s["seriesCd"], {"name": s["name"], "vals": {}})
            for x, v in zip(xs, s["data"]):
                ent["vals"][int(x)] = v
    for cc, ent in sorted(series.items(), key=lambda kv: -sum(kv[1]["vals"].values())):
        put(ws, r, 1, kor); put(ws, r, 2, sym); put(ws, r, 3, cc); put(ws, r, 4, ent["name"])
        for j, y in enumerate(YEARS):
            put(ws, r, 5 + j, ent["vals"].get(y, 0), FMT_INT)
        put(ws, r, 5 + len(YEARS), f"=SUM({get_column_letter(5)}{r}:{get_column_letter(4+len(YEARS))}{r})", FMT_INT, fill=fill_formula)
        r += 1
finish(ws, UNB_FIRST - 1, r - 1, len(hdr))

# 12) UN_한국수출
S_UNX = "UN_한국→수입국_연도별(수출)"
hdr = ["광종", "원소", "연도", "순위", "수입국코드", "수입국명", "금액(USD)", "금액(천USD)", "중량(kg)", "중량(톤)", "금액 비중(%, KOMIS)", "중량 비중(%, KOMIS)", "한국 총수출금액(USD)", "한국 총수출중량(kg)"]
ws, r = new_sheet(S_UNX, hdr, [8, 6, 6, 5, 8, 12, 15, 13, 14, 12, 12, 12, 16, 16], note="출처: KOMIS 글로벌 광물지도 (UN Comtrade), 수출국=대한민국. " + UN_NOTE.split("주의:")[1].strip())
UNX_FIRST = r
for code, sym, kor, eng in METALS:
    for y in YEARS:
        d = load(f"N_list_krexp_{code}_{y}")
        for row in (d or {}).get("list", []):
            put(ws, r, 1, kor); put(ws, r, 2, sym); put(ws, r, 3, y); put(ws, r, 4, row["RNUM"]); put(ws, r, 5, row["incmNtnCd"]); put(ws, r, 6, row["incmNtnNm"])
            put(ws, r, 7, row["amt"], FMT_INT); put(ws, r, 8, f"=G{r}/1000", FMT_1, fill=fill_formula)
            put(ws, r, 9, row["weig"], FMT_2); put(ws, r, 10, f"=I{r}/1000", FMT_2, fill=fill_formula)
            put(ws, r, 11, float(row["amtRate"]) / 100, FMT_PCT); put(ws, r, 12, float(row["weigRate"]) / 100, FMT_PCT)
            put(ws, r, 13, row["sumAmt"], FMT_INT); put(ws, r, 14, row["sumWeig"], FMT_2)
            r += 1
finish(ws, UNX_FIRST - 1, r - 1, len(hdr))

# 13) UN_의존도
S_UND = "UN_의존도(수출국비중)"
hdr = ["광종", "원소", "연도", "순위", "수출국코드", "수출국명", "→한국 금액(USD)", "→한국 중량(kg)",
       "수출국 전세계 수출총액(USD)", "수출국 수출 중 한국 비중(%)", "수출국 전세계 수출총중량(kg)", "수출국 수출중량 중 한국 비중(%)",
       "한국 총수입금액(USD)", "한국 수입 중 해당국 비중(%)", "한국 총수입중량(kg)", "한국 수입중량 중 해당국 비중(%)"]
ws, r = new_sheet(S_UND, hdr, [8, 6, 6, 5, 8, 12, 15, 14, 18, 14, 18, 14, 16, 14, 16, 14],
                  note="출처: KOMIS 글로벌 광물지도 지도 API(getListMapNationData, 수입국=대한민국). 수출국 입장에서 전체 수출 중 한국향 비중과, 한국 수입 중 해당국 비중을 함께 제공. " + UN_NOTE.split("주의:")[1].strip())
UND_FIRST = r
for code, sym, kor, eng in METALS:
    for y in YEARS:
        d = load(f"N_map_krimp_{code}_{y}")
        for row in sorted((d or {}).get("data", {}).get("mapData", []), key=lambda x: x.get("RNUM", 0)):
            put(ws, r, 1, kor); put(ws, r, 2, sym); put(ws, r, 3, y); put(ws, r, 4, row.get("RNUM")); put(ws, r, 5, row["crtrNtnCd"]); put(ws, r, 6, row["crtrNtnKornNm"])
            put(ws, r, 7, row["price"], FMT_INT); put(ws, r, 8, row["weight"], FMT_2)
            put(ws, r, 9, row["crtrTotalAmt"], FMT_INT); put(ws, r, 10, float(row["crtrNtnAmtRt"]) / 100, FMT_PCT)
            put(ws, r, 11, row["crtrTotalWeig"], FMT_2); put(ws, r, 12, float(row["crtrNtnWeigRt"]) / 100, FMT_PCT)
            put(ws, r, 13, row["trgtTotalAmt"], FMT_INT); put(ws, r, 14, float(row["trgtNtnAmtRt"]) / 100, FMT_PCT)
            put(ws, r, 15, row["trgtTotalWeig"], FMT_2); put(ws, r, 16, float(row["trgtNtnWeigRt"]) / 100, FMT_PCT)
            r += 1
finish(ws, UND_FIRST - 1, r - 1, len(hdr))

# 14) UN_글로벌 교역흐름
S_UNG = "UN_글로벌_교역흐름TOP30"
hdr = ["광종", "원소", "연도", "순위", "수출국코드", "수출국", "수입국코드", "수입국", "금액(USD)", "금액(천USD)", "중량(kg)", "중량(톤)", "금액 비중(%, 전세계)", "중량 비중(%, 전세계)", "전세계 교역총액(USD)", "전세계 교역총중량(kg)"]
ws, r = new_sheet(S_UNG, hdr, [8, 6, 6, 5, 8, 12, 8, 12, 16, 14, 15, 12, 12, 12, 18, 18],
                  note="출처: KOMIS 글로벌 광물지도 세부현황(UN Comtrade), 국가 미지정 시 전세계 수출국→수입국 교역흐름 상위 30건(API가 30건으로 제한). 2025~2026년은 일부 보고국만 반영된 부분 집계. " + UN_NOTE.split("주의:")[1].strip())
UNG_FIRST = r
for code, sym, kor, eng in METALS:
    for y in YEARS:
        d = load(f"N_list_global_{code}_{y}")
        for row in (d or {}).get("list", []):
            put(ws, r, 1, kor); put(ws, r, 2, sym); put(ws, r, 3, y); put(ws, r, 4, row["RNUM"]); put(ws, r, 5, row["expNtnCd"]); put(ws, r, 6, row["expNtnNm"]); put(ws, r, 7, row["incmNtnCd"]); put(ws, r, 8, row["incmNtnNm"])
            put(ws, r, 9, row["amt"], FMT_INT); put(ws, r, 10, f"=I{r}/1000", FMT_1, fill=fill_formula)
            put(ws, r, 11, row["weig"], FMT_2); put(ws, r, 12, f"=K{r}/1000", FMT_2, fill=fill_formula)
            put(ws, r, 13, float(row["amtRate"]) / 100, FMT_PCT); put(ws, r, 14, float(row["weigRate"]) / 100, FMT_PCT)
            put(ws, r, 15, row["sumAmt"], FMT_INT); put(ws, r, 16, row["sumWeig"], FMT_2)
            r += 1
finish(ws, UNG_FIRST - 1, r - 1, len(hdr))

# 15) UN_글로벌 국가별 추이
S_UNGB = "UN_글로벌_국가별추이"
hdr = ["광종", "원소", "구분", "국가코드", "국가명"] + [f"{y} 금액(USD)" for y in YEARS] + ["합계(USD)"]
ws, r = new_sheet(S_UNGB, hdr, [8, 6, 8, 8, 12] + [16] * len(YEARS) + [16],
                  note="출처: KOMIS 글로벌 광물지도 '국가별 교역규모 Top10' 차트 API(국가 미지정). 구분 '수입국'=전세계 주요 수입국별 수입액, '수출국'=주요 수출국별 수출액. UN Comtrade. 2025~2026년은 일부 보고국만 반영된 부분 집계(한국 미반영). " + UN_NOTE.split("주의:")[1].strip())
UNGB_FIRST = r
for code, sym, kor, eng in METALS:
    for ie, lab in (("I", "수입국"), ("E", "수출국")):
        series = {}
        for ye in (2024, 2026):
            d = load(f"N_bar_global{ie}_{code}_{ye}")
            bc = (d or {}).get("data", {}).get("barChart", {})
            xs = bc.get("xaxis", [])
            for s in bc.get("series", []):
                ent = series.setdefault(s["seriesCd"], {"name": s["name"], "vals": {}})
                for x, v in zip(xs, s["data"]):
                    ent["vals"][int(x)] = v
        for cc, ent in sorted(series.items(), key=lambda kv: -sum(kv[1]["vals"].values())):
            put(ws, r, 1, kor); put(ws, r, 2, sym); put(ws, r, 3, lab); put(ws, r, 4, cc); put(ws, r, 5, ent["name"])
            for j, y in enumerate(YEARS):
                put(ws, r, 6 + j, ent["vals"].get(y, 0), FMT_INT)
            put(ws, r, 6 + len(YEARS), f"=SUM({get_column_letter(6)}{r}:{get_column_letter(5+len(YEARS))}{r})", FMT_INT, fill=fill_formula)
            r += 1
finish(ws, UNGB_FIRST - 1, r - 1, len(hdr))

# 16) 코드표
S_CD = "코드표"
ws = wb.create_sheet(S_CD)
rr = 1
ws.cell(row=rr, column=1, value="광종 코드 (KOMIS)").font = f_bold; rr += 1
for i, h in enumerate(["KOMIS 광종코드", "광종(한글)", "원소", "영문"], 1):
    put(ws, rr, i, h, bold=True, fill=fill_sub)
rr += 1
for code, sym, kor, eng in METALS:
    put(ws, rr, 1, code); put(ws, rr, 2, kor); put(ws, rr, 3, sym); put(ws, rr, 4, eng); rr += 1
rr += 1
ws.cell(row=rr, column=1, value="관세청 HS코드(10자리) - 대한민국 광물지도").font = f_bold; rr += 1
for i, h in enumerate(["광종", "원소", "HS코드", "품목명"], 1):
    put(ws, rr, i, h, bold=True, fill=fill_sub)
rr += 1
for code, sym, kor, eng in METALS:
    for h in HS_KR[code]:
        put(ws, rr, 1, kor); put(ws, rr, 2, sym); put(ws, rr, 3, h["hsCd"]); put(ws, rr, 4, h["itemNm"]); rr += 1
rr += 1
ws.cell(row=rr, column=1, value="UN HS코드(6자리) - 글로벌 광물지도 (UN Comtrade)").font = f_bold; rr += 1
for i, h in enumerate(["광종", "원소", "HS6", "매핑된 관세청 HS코드", "품목(영문)"], 1):
    put(ws, rr, i, h, bold=True, fill=fill_sub)
rr += 1
for code, sym, kor, eng in METALS:
    for h in HS_UN[code]:
        put(ws, rr, 1, kor); put(ws, rr, 2, sym); put(ws, rr, 3, h["hsCdUn"]); put(ws, rr, 4, h["hsCd"]); put(ws, rr, 5, h["rmrk"]); rr += 1
rr += 1
ws.cell(row=rr, column=1, value="생산품 유형(물질흐름) 코드").font = f_bold; rr += 1
for i, h in enumerate(["유형코드", "유형명"], 1):
    put(ws, rr, i, h, bold=True, fill=fill_sub)
rr += 1
for k, v in MF_NM.items():
    put(ws, rr, 1, k); put(ws, rr, 2, v); rr += 1
rr += 1
ws.cell(row=rr, column=1, value="국가코드 (본 파일에 등장하는 국가)").font = f_bold; rr += 1
for i, h in enumerate(["국가코드", "국가명(한글)", "국가명(영문)", "권역"], 1):
    put(ws, rr, i, h, bold=True, fill=fill_sub)
rr += 1
used = set()
for f in glob.glob(os.path.join(D, "*.json")):
    j = json.load(open(f, encoding="utf-8"))["resp"]
    for row in j.get("list", []) if isinstance(j, dict) else []:
        for k in ("ntnCd", "expNtnCd", "incmNtnCd"):
            if row.get(k): used.add(row[k])
for cc in sorted(used):
    put(ws, rr, 1, cc); put(ws, rr, 2, NTN.get(cc, ("",))[0]); put(ws, rr, 3, ENG.get(cc, "")); put(ws, rr, 4, NTN.get(cc, ("", ""))[1]); rr += 1
for i, w in enumerate([16, 14, 14, 40, 60], 1):
    ws.column_dimensions[get_column_letter(i)].width = w

# =====================================================================
# 0) 개요 sheet (first)
# =====================================================================
ws = wb.create_sheet("개요", 0)
ws.column_dimensions["A"].width = 28; ws.column_dimensions["B"].width = 120
lines = [
    ("KOMIS 희소금속(Ta·Nb·In·Ga·Hf·Re) 국내 수입현황 2020~2026", None),
    ("", None),
    ("추출일", datetime.date.today().strftime("%Y-%m-%d")),
    ("출처 1 (국내)", "KOMIS 핵심광물지도 > 대한민국  https://www.komis.or.kr/Komis/MnrlMap/Korea  (자료출처: 관세청 수출입통계, 공공데이터포털)"),
    ("출처 2 (글로벌)", "KOMIS 핵심광물지도 > 글로벌  https://www.komis.or.kr/Komis/MnrlMap/Nation  (자료출처: UN Comtrade)"),
    ("대상 광종", "탄탈륨(Ta, MNRL0026), 니오븀(Nb, MNRL0007), 인듐(In, MNRL0025), 갈륨(Ga, MNRL0024), 하프늄(Hf, MNRL0040), 레늄(Re, MNRL0042)"),
    ("대상 기간", "2020년~2026년 (연도별). 관세청 데이터의 2026년은 1~9월 누계(KOMIS 기준연도 목록상 2026년 9월까지 집계). 월별 데이터는 2020.01~2026.09."),
    ("단위", "원자료는 금액 USD, 중량 kg (KOMIS 화면의 천USD·톤 표시는 /1000 환산). 본 파일은 USD·kg 원자료 열과 천USD·톤 수식 열을 함께 제공."),
    ("비중(점유율)", "관세청 시트: 해당 광종·연도 총수입금액(중량) 대비 수식(SUMIFS). UN 시트: KOMIS가 제공한 비중(%) 값을 그대로 수록(소수 → 백분율 서식)."),
    ("", None),
    ("시트 구성", None),
    ("요약", "광종×연도별 총수입금액·중량, 전년비, 수출, 수입국 수, 상위 3개 수입국과 금액비중 (전부 수식, '관세청_연도별_국가' 참조)"),
    ("요약_매트릭스", "광종(행)×연도(열) 매트릭스: 수입금액, 수입중량, 수출금액, 수출중량, 전년비"),
    ("관세청_연도별_국가", "[핵심] 광종·연도·국가별 수입/수출 금액·중량, 비중(수식). 모든 수입국 전체(TOP30 제한 없음)"),
    ("관세청_연도별_총계(API)", "KOMIS '수입규모 추이' 차트 API가 돌려준 연도별 총계 (요약 시트 SUMIFS와 교차검증용, 값 일치 확인)"),
    ("관세청_HS코드별_국가 / _총계", "관세청 HS 10자리 코드별(광석·괴·가루·스크랩·기타 등) 국가별 수입/수출 및 연도별 총계 매트릭스"),
    ("관세청_생산품유형별_국가 / _총계", "KOMIS 물질흐름 분류(기초원료/기초금속/가공금속/스크랩)별 국가별 수입/수출 및 연도별 총계"),
    ("관세청_월별_국가 / _총계", "2020.01~2026.09 월별 국가별 수입/수출, 광종별 월별 총계 매트릭스"),
    ("관세청_전년대비증가TOP5", "KOMIS '전년대비 교역증가 TOP5' API 결과 (최근 5개년 2021~2025 수입금액, 2025년 증가율). 조회연도와 무관하게 동일값 반환됨"),
    ("UN_수출국→한국_연도별", "UN Comtrade 기준 한국 수입(수출국별) 금액·중량·비중 (2020~2024, 2025 이후 미집계)"),
    ("UN_HS6별_수출국→한국", "UN HS 6자리별 수출국→한국 수입 현황"),
    ("UN_수출국별_추이(한국수입)", "수출국별 연도별 한국향 수출액 추이 (KOMIS 차트 API)"),
    ("UN_한국→수입국_연도별(수출)", "UN Comtrade 기준 한국의 수출(수입국별)"),
    ("UN_의존도(수출국비중)", "수출국의 전세계 수출 중 한국향 비중 + 한국 수입 중 해당국 비중 (상호 의존도)"),
    ("UN_글로벌_교역흐름TOP30", "전세계 수출국→수입국 교역흐름 상위 30건 (광종·연도별, 2025~2026은 부분 집계)"),
    ("UN_글로벌_국가별추이", "전세계 주요 수입국/수출국별 연도별 교역액 추이 (2025~2026은 부분 집계)"),
    ("코드표", "광종코드, 관세청 HS10 코드, UN HS6 코드, 생산품 유형, 국가코드"),
    ("", None),
    ("주의사항", None),
    ("1", "관세청(10자리 HS) 데이터는 금속별로 정확히 구분되나, UN Comtrade(6자리 HS)는 811292/811299가 갈륨·게르마늄·인듐·니오븀·바나듐 공용, 261590이 니오븀·탄탈륨·바나듐 광석 공용이므로 UN 시트의 인듐·갈륨 수치는 동일하게 집계되고 니오븀도 타 금속이 섞여 있음. 금속별 분석은 관세청 시트를 사용할 것."),
    ("2", "관세청 데이터와 UN Comtrade 데이터는 집계 기준(신고국, 환율, 시점)이 달라 소폭 차이가 있음 (예: 탄탈륨 2024 중국 38,126,993 vs 38,155,370 USD)."),
    ("3", "KOMIS 글로벌(UN) 데이터는 2025년·2026년이 아직 집계되지 않아 해당 연도 행은 비어 있음."),
    ("4", "'관세청_전년대비증가TOP5'는 KOMIS 화면에서 숨겨진 API로, 검색연도를 바꿔도 최근 5개년 고정값이 반환됨(API 특성)."),
    ("5", "수식 셀은 회색 배경. 원자료 셀은 흰색 배경. 비중은 소수(0.4134)로 저장되고 % 서식으로 표시됨."),
    ("6", "하프늄·레늄은 2020~2021년 데이터가 관세청·UN 모두 없음(0). HS 2022 개정으로 신설된 전용 코드(8112.31/39, 8112.41/49)만 KOMIS에 매핑되어 있고, 2021년 이전에는 8112.92/99 '기타'에 포함되어 분리 집계가 불가능함."),
    ("7", "UN 글로벌 시트의 2025~2026년 값은 일부 보고국만 반영된 부분 집계이며 한국은 미반영. 국가별 추이 비교는 2024년까지만 유효."),
    ("8", "원자료(JSON)는 KOMIS 화면이 호출하는 내부 AJAX API(getListKoreaData, getLineChartDataKorea, getTop5TableDataKorea, getListMapKoreaData, getListDataNation, getBarChartDataNation, getListMapNationData)를 그대로 호출해 수집함. 화면의 TOP30 제한 없이 전체 국가를 수록."),
]
for i, (a, b) in enumerate(lines, 1):
    ca = ws.cell(row=i, column=1, value=a)
    ca.font = f_title if i == 1 else (f_bold if b is None and a else f_norm)
    if b is not None:
        cb = ws.cell(row=i, column=2, value=b); cb.font = f_norm; cb.alignment = Alignment(wrap_text=True, vertical="top")
ws.freeze_panes = None

wb.save(OUT)
print("saved", OUT)
print("rows: KY", KY_LAST - KY_FIRST + 1, "HS", HS_LAST - HS_FIRST + 1, "MO", MO_LAST - MO_FIRST + 1)
